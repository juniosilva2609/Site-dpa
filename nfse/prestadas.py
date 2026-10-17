"""NFS-e prestadas pela C3S no Portal Nacional -> planilha de importação do Alterdata.

Lê a distribuição de DF-e do ADN (mTLS com o e-CNPJ A1 da C3S), fica só com as
NFS-e em que a C3S é a prestadora, emitidas no mês pedido e ainda válidas
(sem evento de cancelamento nem de substituição), e gera o Excel:

    DATA EMISSAO | NUM NFSE | VALOR | CLIENTE | HISTORICO CONTABIL | CONTA DEBITO | CONTA CREDITO

Cada nota vira uma linha com o valor bruto (débito = conta do cliente, crédito
= 861). Quando o ISS foi retido pelo tomador, vem logo abaixo uma segunda linha
com só o valor do ISS (débito 1393, crédito = conta do cliente).

Uso:  python -m nfse.prestadas AAAA-MM <pasta de saída> [c3s|licitprint]
Credenciais (variáveis de ambiente, mesmo formato dos bancos): DPA_C3S_NFSE_CERT e
DPA_C3S_NFSE_KEY (PEM ou PEM em base64, ver connectors/base.py).
"""

from __future__ import annotations

import base64
import gzip
import html
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path

CNPJ_C3S = "47553447000151"
# Uma entrada por empresa: CNPJ do prestador, variaveis com o e-CNPJ (PEM) e mapa CNPJ do cliente -> conta
EMPRESAS = {
    "c3s": {"cnpj": CNPJ_C3S, "cert": "DPA_C3S_NFSE_CERT", "key": "DPA_C3S_NFSE_KEY", "contas": "contas_clientes_c3s.json"},
    "licitprint": {"cnpj": "65942303000179", "cert": "DPA_LICITPRINT_NFSE_CERT", "key": "DPA_LICITPRINT_NFSE_KEY",
                   "contas": "contas_clientes_licitprint.json"},
}
ADN = "https://adn.nfse.gov.br/contribuintes/DFe/{nsu}"
CONTA_RECEITA = 861
CONTA_ISS_RETIDO = 1393
# Eventos que tiram a nota da base: cancelamento, cancelamento por substituição
# (o evento aponta a nota SUBSTITUÍDA) e cancelamento deferido por análise fiscal.
EVENTOS_QUE_CANCELAM = {"e101101", "e105102", "e105104"}


@dataclass
class Nota:
    chave: str
    numero: int
    emissao: date
    valor: float
    iss: float
    iss_retido: bool
    cnpj_cliente: str
    nome_xml: str


def _tag(xml: str, nome: str) -> str | None:
    m = re.search(rf"<{nome}>([^<]*)<", xml)
    return html.unescape(m.group(1)) if m else None


def baixar_dfe(cert: tuple[str, str], cnpj: str = CNPJ_C3S) -> list[dict]:
    import requests

    docs, nsu = [], 0
    while True:
        r = requests.get(ADN.format(nsu=nsu), params={"cnpjConsulta": cnpj, "lote": "true"}, cert=cert, timeout=90)
        if r.status_code == 404:  # E2220: nenhum documento além deste NSU
            break
        r.raise_for_status()
        lote = r.json().get("LoteDFe") or []
        if not lote:
            break
        for d in lote:
            d["xml"] = gzip.decompress(base64.b64decode(d.pop("ArquivoXml"))).decode("utf-8")
            docs.append(d)
        nsu = max(int(d["NSU"]) for d in lote)
    return docs


def canceladas(docs: list[dict]) -> set[str]:
    """Chaves de NFS-e alvo de evento que a invalida."""
    chaves = set()
    for d in docs:
        if d["TipoDocumento"] != "EVENTO":
            continue
        tipo = re.search(r"<(e\d{6})>", d["xml"])
        chave = _tag(d["xml"], "chNFSe")
        if tipo and chave and tipo.group(1) in EVENTOS_QUE_CANCELAM:
            chaves.add(chave)
    return chaves


def notas_validas(docs: list[dict], ano: int, mes: int, cnpj: str = CNPJ_C3S) -> list[Nota]:
    fora = canceladas(docs)
    notas = []
    for d in docs:
        if d["TipoDocumento"] != "NFSE" or d["ChaveAcesso"] in fora:
            continue
        x = d["xml"]
        emit = re.search(r"<emit><CNPJ>(\d+)</CNPJ>", x)
        dh = _tag(x, "dhEmi") or ""
        if not emit or emit.group(1) != cnpj or dh[:7] != f"{ano:04d}-{mes:02d}" or _tag(x, "cStat") != "100":
            continue
        m_toma = re.search(r"<toma>.*?</toma>", x, re.S)   # tomador pode nao ser identificado na nota
        toma = m_toma.group(0) if m_toma else ""
        doc = re.search(r"<(?:CNPJ|CPF)>(\d+)<", toma)
        notas.append(Nota(
            chave=d["ChaveAcesso"], numero=int(_tag(x, "nNFSe")), emissao=date.fromisoformat(dh[:10]),
            valor=float(_tag(x, "vServ")), iss=float(_tag(x, "vISSQN") or 0), iss_retido=_tag(x, "tpRetISSQN") == "2",
            cnpj_cliente=doc.group(1) if doc else "", nome_xml=_tag(toma, "xNome") or "TOMADOR NAO IDENTIFICADO"))
    return sorted(notas, key=lambda n: (n.emissao, n.numero))


def linhas(notas: list[Nota], contas: dict[str, dict]) -> list[tuple]:
    out = []
    for n in notas:
        c = contas.get(n.cnpj_cliente)
        nome = (c["nome"] if c else n.nome_xml).upper()
        conta = c["conta"] if c else None
        out.append((n.emissao, n.numero, n.valor, nome, f"PREST. SERV. CONF. NFSE {n.numero} CLIENTE {nome}", conta, CONTA_RECEITA))
        if n.iss_retido and n.iss:
            out.append((n.emissao, n.numero, n.iss, nome, f"ISS RETIDO SERV. PREST. CONF. NFSE {n.numero} CLIENTE {nome}", CONTA_ISS_RETIDO, conta))
    return out


def gerar_excel(linhas_: list[tuple], destino: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Importação"
    ws.append(["DATA EMISSAO", "NUM NFSE", "VALOR", "CLIENTE", "HISTORICO CONTABIL", "CONTA DEBITO", "CONTA CREDITO"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center")
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for i, l in enumerate(linhas_, start=2):
        ws.append(list(l))
        ws.cell(i, 1).number_format = "dd/mm/yyyy"
        ws.cell(i, 3).number_format = "#,##0.00"
        if l[5] is None or l[6] is None:
            for c in ws[i]:
                c.fill = amarelo
    for col, w in zip("ABCDEFG", [14, 11, 14, 52, 72, 14, 15]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(destino)


def main(competencia: str, pasta: str, empresa: str = "c3s") -> None:
    from connectors.base import arquivo_de_credencial

    cfg = EMPRESAS[empresa]
    ano, mes = map(int, competencia.split("-"))
    docs = baixar_dfe((arquivo_de_credencial(cfg["cert"]), arquivo_de_credencial(cfg["key"])), cfg["cnpj"])
    contas_arq = Path(__file__).with_name(cfg["contas"])
    contas = json.loads(contas_arq.read_text(encoding="utf-8")) if contas_arq.exists() else {}
    notas = notas_validas(docs, ano, mes, cfg["cnpj"])
    ls = linhas(notas, contas)
    destino = Path(pasta) / f"IMPORT ALTERDATA NFSE SERV PREST {mes:02d}-{ano}.xlsx"
    gerar_excel(ls, destino)
    fora = canceladas(docs)
    sem_conta = sorted({(n.cnpj_cliente, n.nome_xml) for n in notas if n.cnpj_cliente not in contas})
    print(json.dumps({
        "empresa": empresa, "arquivo": str(destino), "notas": len(notas), "linhas": len(ls),
        "total_bruto": round(sum(n.valor for n in notas), 2), "iss_retido": round(sum(n.iss for n in notas if n.iss_retido), 2),
        "canceladas_ou_substituidas_no_mes": sum(1 for d in docs if d["ChaveAcesso"] in fora and (_tag(d["xml"], "dhEmi") or "")[:7] == f"{ano:04d}-{mes:02d}"),
        "clientes_sem_conta": sem_conta}, ensure_ascii=False))


if __name__ == "__main__":
    main(*sys.argv[1:4])
